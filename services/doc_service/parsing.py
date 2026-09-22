"""Extract plain text from an uploaded application document.

Scope is deliberately narrow: PDF, DOCX, plain text and Markdown cover CVs,
statements of purpose, motivation letters and recommendation letters. Anything
else is rejected at upload with a message that says what to convert it to.
"""

import io
import re
from dataclasses import dataclass, field

from docx import Document as DocxDocument
from pypdf import PdfReader

from shared.errors import ValidationError

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TEXT_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class ParsedDocument:
    text: str
    page_count: int | None = None
    # Character offset where each page starts, so a chunk can be traced back to
    # a page number in the citation the user sees.
    page_offsets: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ")
    # Collapse runs of blank lines but keep paragraph breaks.
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError(f"Could not read that PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ValidationError(
                "That PDF is password-protected. Remove the password and upload it again."
            ) from exc

    parts: list[str] = []
    offsets: list[int] = []
    running = 0
    warnings: list[str] = []

    for index, page in enumerate(reader.pages):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
            warnings.append(f"Page {index + 1} could not be read.")
        cleaned = _clean(page_text)
        offsets.append(running)
        parts.append(cleaned)
        running += len(cleaned) + 2

    text = "\n\n".join(parts).strip()
    if len(text) < 40:
        raise ValidationError(
            "No selectable text was found in that PDF — it looks like a scan. "
            "Upload a text-based PDF or a DOCX, or run OCR on it first."
        )
    return ParsedDocument(
        text=text, page_count=len(reader.pages), page_offsets=offsets, warnings=warnings
    )


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValidationError(f"Could not read that Word document: {exc}") from exc

    blocks = [paragraph.text for paragraph in document.paragraphs]
    # Tables are common in CVs and carry real content, so pull them in too.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = _clean("\n".join(block for block in blocks if block.strip()))
    if len(text) < 40:
        raise ValidationError("That document appears to be empty.")
    return ParsedDocument(text=text)


def _parse_text(data: bytes) -> ParsedDocument:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = _clean(data.decode(encoding))
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValidationError("Could not decode that text file.")

    if len(text) < 40:
        raise ValidationError("That document appears to be empty.")
    return ParsedDocument(text=text)


def parse(data: bytes, content_type: str, filename: str) -> ParsedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if content_type in PDF_TYPES or suffix == ".pdf":
        return _parse_pdf(data)
    if content_type in DOCX_TYPES or suffix == ".docx":
        return _parse_docx(data)
    if content_type in TEXT_TYPES or suffix in {".txt", ".md"}:
        return _parse_text(data)

    if suffix == ".doc":
        raise ValidationError(
            "Legacy .doc files are not supported. Save it as .docx or export a PDF."
        )
    raise ValidationError(
        f"Unsupported file type '{content_type or suffix or 'unknown'}'. "
        "Upload a PDF, DOCX, TXT or MD file."
    )
